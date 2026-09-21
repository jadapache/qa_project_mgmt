import io
import uuid
import pytest
import openpyxl
from fastapi import HTTPException
from sqlalchemy.ext.asyncio import create_async_engine, AsyncSession
from sqlalchemy.orm import sessionmaker
from app.models.models import Base, TestCase as TestCaseModel
from app.services.import_service import validate_and_import_xlsx

TEST_DATABASE_URL = "sqlite+aiosqlite:///:memory:"

@pytest.mark.asyncio
async def test_import_invalid_file_format():
    with pytest.raises(HTTPException) as exc_info:
        await validate_and_import_xlsx(
            file_bytes=b"invalid corrupt bytes",
            entity="stories",
            project_id=str(uuid.uuid4()),
            iteration_id=str(uuid.uuid4()),
            db=None
        )
    assert exc_info.value.status_code == 422
    assert "no es una hoja de cálculo XLSX válida" in exc_info.value.detail

@pytest.mark.asyncio
async def test_import_empty_file():
    wb = openpyxl.Workbook()
    output = io.BytesIO()
    wb.save(output)
    
    with pytest.raises(HTTPException) as exc_info:
        await validate_and_import_xlsx(
            file_bytes=output.getvalue(),
            entity="stories",
            project_id=str(uuid.uuid4()),
            iteration_id=str(uuid.uuid4()),
            db=None
        )
    assert exc_info.value.status_code == 422
    assert "está vacío" in exc_info.value.detail

@pytest.mark.asyncio
async def test_import_unsupported_entity():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Col1", "Col2"])
    ws.append(["Val1", "Val2"])
    output = io.BytesIO()
    wb.save(output)

    with pytest.raises(HTTPException) as exc_info:
        await validate_and_import_xlsx(
            file_bytes=output.getvalue(),
            entity="unsupported_entity",
            project_id=str(uuid.uuid4()),
            iteration_id=str(uuid.uuid4()),
            db=None
        )
    assert exc_info.value.status_code == 422
    assert "no soportada" in exc_info.value.detail

@pytest.mark.asyncio
async def test_import_atomicity_rejects_invalid_row():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Descripción", "Criterios", "Prioridad"])
    ws.append(["Historia Valida 1", "Criterio A", "Alta"])
    ws.append(["", "Criterio " + "x"*2001, "InvalidPriority"])

    output = io.BytesIO()
    wb.save(output)

    with pytest.raises(HTTPException) as exc_info:
        await validate_and_import_xlsx(
            file_bytes=output.getvalue(),
            entity="stories",
            project_id=str(uuid.uuid4()),
            iteration_id=str(uuid.uuid4()),
            db=None
        )
    assert exc_info.value.status_code == 422
    detail = exc_info.value.detail
    assert "row_errors" in detail
    assert detail["total_errors"] > 0

@pytest.mark.asyncio
async def test_import_test_cases_success():
    engine = create_async_engine(TEST_DATABASE_URL, echo=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    async_session = sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Título", "Precondiciones", "Datos Entrada", "Resultado Esperado"])
    ws.append(["Caso de Prueba Importado 1", "Precond 1", "Datos 1", "Esperado 1"])
    ws.append(["Title " + "x"*300, "", "", ""]) # Invalid row (>255 chars)

    output = io.BytesIO()
    wb.save(output)

    async with async_session() as session:
        # Invalid row test
        with pytest.raises(HTTPException) as exc_info:
            await validate_and_import_xlsx(
                file_bytes=output.getvalue(),
                entity="test_cases",
                project_id=str(uuid.uuid4()),
                iteration_id=str(uuid.uuid4()),
                db=session
            )
        assert exc_info.value.status_code == 422

        # Valid import test
        wb_valid = openpyxl.Workbook()
        ws_v = wb_valid.active
        ws_v.append(["Título", "Precondiciones", "Datos Entrada", "Resultado Esperado"])
        ws_v.append(["Caso Válido 1", "Pre1", "Data1", "Exp1"])
        output_valid = io.BytesIO()
        wb_valid.save(output_valid)

        count, errs = await validate_and_import_xlsx(
            file_bytes=output_valid.getvalue(),
            entity="test_cases",
            project_id=str(uuid.uuid4()),
            iteration_id=str(uuid.uuid4()),
            db=session
        )
        assert count == 1
        assert errs == []

    await engine.dispose()
