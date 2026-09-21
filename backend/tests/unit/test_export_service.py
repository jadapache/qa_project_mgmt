import pytest
from datetime import date
import openpyxl
from sqlalchemy.ext.asyncio import create_async_engine, AsyncSession
from sqlalchemy.orm import sessionmaker
from app.models.models import Base, User, TestCase as TestCaseModel, UserStory, Defect, Project
from app.services.export_service import generate_xlsx

TEST_DATABASE_URL = "sqlite+aiosqlite:///:memory:"

@pytest.mark.asyncio
async def test_generate_xlsx_export():
    engine = create_async_engine(TEST_DATABASE_URL, echo=False)
    async with engine.begin() as conn:
        await conn.run_sync(Base.metadata.create_all)

    async_session = sessionmaker(engine, class_=AsyncSession, expire_on_commit=False)

    async with async_session() as session:
        owner = User(username="owner1", password_hash="hash", role="Administrador")
        session.add(owner)
        await session.flush()

        project = Project(
            name="Project Test",
            start_date=date.today(),
            end_date_estimated=date.today(),
            owner_id=owner.id
        )
        session.add(project)
        await session.flush()

        # Add TestCase
        tc = TestCaseModel(
            story_id=owner.id, # dummy uuid
            title="TestCase 1",
            preconditions="Precondición",
            input_data="Datos",
            expected_result="Esperado",
            status="Diseño"
        )
        # Add UserStory
        us = UserStory(
            iteration_id=owner.id,
            description="Descripción historia",
            acceptance_criteria="Criterio A",
            priority="Alta",
            status="Nuevo"
        )
        # Add Defect
        df = Defect(
            execution_id=owner.id,
            title="Defecto 1",
            description="Descripción defecto",
            steps_to_reproduce="Paso 1",
            severity="Alta",
            status="Nuevo",
            jira_issue_key="JIRA-101"
        )
        session.add_all([tc, us, df])
        await session.commit()

        # Export test cases
        tc_stream = await generate_xlsx(str(project.id), "test_cases", session)
        wb_tc = openpyxl.load_workbook(tc_stream)
        assert "Casos de Prueba" in wb_tc.sheetnames
        ws_tc = wb_tc["Casos de Prueba"]
        assert ws_tc.cell(row=2, column=2).value == "TestCase 1"

        # Export stories
        us_stream = await generate_xlsx(str(project.id), "stories", session)
        wb_us = openpyxl.load_workbook(us_stream)
        assert "Historias de Usuario" in wb_us.sheetnames
        ws_us = wb_us["Historias de Usuario"]
        assert ws_us.cell(row=2, column=2).value == "Descripción historia"

        # Export defects
        df_stream = await generate_xlsx(str(project.id), "defects", session)
        wb_df = openpyxl.load_workbook(df_stream)
        assert "Defectos" in wb_df.sheetnames
        ws_df = wb_df["Defectos"]
        assert ws_df.cell(row=2, column=2).value == "Defecto 1"

        # Export unknown entity
        unk_stream = await generate_xlsx(str(project.id), "unknown", session)
        wb_unk = openpyxl.load_workbook(unk_stream)
        assert "Exportación" in wb_unk.sheetnames

    await engine.dispose()
