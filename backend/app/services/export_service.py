import io
from typing import Any, List
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select

from app.models.models import TestCase, UserStory, Defect, Execution

async def generate_xlsx(project_id: str, entity: str, db: AsyncSession) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.views.sheetView[0].showGridLines = True

    # Estilos corporativos
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid") # Dark Navy
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    data_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    if entity == "test_cases":
        ws.title = "Casos de Prueba"
        headers = ["ID", "Título", "Precondiciones", "Datos Entrada", "Resultado Esperado", "Estado", "Fecha Creación"]
        ws.append(headers)

        result = await db.execute(select(TestCase).order_by(TestCase.created_at.desc()))
        cases = result.scalars().all()
        for c in cases:
            ws.append([
                str(c.id), c.title, c.preconditions or "", c.input_data or "",
                c.expected_result or "", c.status, c.created_at.strftime("%Y-%m-%d %H:%M:%S")
            ])

    elif entity == "stories":
        ws.title = "Historias de Usuario"
        headers = ["ID", "Descripción", "Criterios Aceptación", "Prioridad", "Estado", "Fecha Creación"]
        ws.append(headers)

        result = await db.execute(select(UserStory).order_by(UserStory.created_at.desc()))
        stories = result.scalars().all()
        for s in stories:
            ws.append([
                str(s.id), s.description, s.acceptance_criteria or "",
                s.priority, s.status, s.created_at.strftime("%Y-%m-%d %H:%M:%S")
            ])

    elif entity == "defects":
        ws.title = "Defectos"
        headers = ["ID", "Título", "Descripción", "Pasos Reproducir", "Severidad", "Estado", "Jira Key", "Fecha Creación"]
        ws.append(headers)

        result = await db.execute(select(Defect).order_by(Defect.created_at.desc()))
        defects = result.scalars().all()
        for d in defects:
            ws.append([
                str(d.id), d.title, d.description or "", d.steps_to_reproduce or "",
                d.severity, d.status, d.jira_issue_key or "", d.created_at.strftime("%Y-%m-%d %H:%M:%S")
            ])

    else:
        ws.title = "Exportación"
        ws.append(["ID", "Entidad Desconocida"])

    # Estilar encabezados
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Ajustar ancho de columnas
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output
