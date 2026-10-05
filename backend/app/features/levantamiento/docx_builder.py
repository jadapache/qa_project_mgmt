from __future__ import annotations

import io
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

PRIMARY_COLOR = RGBColor(0, 39, 119)    # #002777 Corporate Navy
SECONDARY_COLOR = RGBColor(0, 68, 151)  # #004497 Navy Accent
ACCENT_GREEN = RGBColor(5, 150, 105)    # #059669 Green
TEXT_COLOR = RGBColor(51, 65, 85)       # #334155 Dark Gray
BG_HEADER_HEX = "002777"
BG_ALT_ROW_HEX = "F8FAFC"


def set_cell_background(cell, hex_color: str):
  tcPr = cell._element.get_or_add_tcPr()
  shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
  tcPr.append(shd)


def create_levantamiento_docx(markdown_content: str, title: str = "Levantamiento Detallado de Requerimientos y Especificación Funcional") -> bytes:
  doc = Document()

  for section in doc.sections:
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

  # Title Header
  p_title = doc.add_paragraph()
  p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
  p_title.paragraph_format.space_before = Pt(0)
  p_title.paragraph_format.space_after = Pt(4)
  run_title = p_title.add_run(title.upper())
  run_title.font.name = "Segoe UI"
  run_title.font.size = Pt(18)
  run_title.font.bold = True
  run_title.font.color.rgb = PRIMARY_COLOR

  p_sub = doc.add_paragraph()
  p_sub.paragraph_format.space_after = Pt(16)
  run_sub = p_sub.add_run("Documentación Funcional y Criterios de Aceptación - QA Project MGMT")
  run_sub.font.name = "Segoe UI"
  run_sub.font.size = Pt(9.5)
  run_sub.font.italic = True
  run_sub.font.color.rgb = RGBColor(100, 116, 139)

  lines = markdown_content.splitlines()
  in_code_block = False
  table_rows: list[list[str]] = []
  in_table = False

  for line in lines:
    stripped = line.strip()

    if stripped.startswith("```"):
      in_code_block = not in_code_block
      continue

    if in_code_block:
      p = doc.add_paragraph()
      p.paragraph_format.left_indent = Inches(0.4)
      p.paragraph_format.space_after = Pt(2)
      run = p.add_run(line)
      run.font.name = "Consolas"
      run.font.size = Pt(9.5)
      run.font.color.rgb = RGBColor(71, 85, 105)
      continue

    # Table processing
    if stripped.startswith("|") and stripped.endswith("|"):
      if re.match(r"^\|[\s\-:|]+\|$", stripped):
        continue
      cells = [c.strip() for c in stripped.split("|")[1:-1]]
      table_rows.append(cells)
      in_table = True
      continue
    elif in_table:
      _render_table(doc, table_rows)
      table_rows = []
      in_table = False

    if not stripped:
      continue

    # Headings
    if stripped.startswith("# "):
      p = doc.add_paragraph()
      p.paragraph_format.space_before = Pt(16)
      p.paragraph_format.space_after = Pt(6)
      run = p.add_run(stripped[2:])
      run.font.name = "Segoe UI"
      run.font.size = Pt(15)
      run.font.bold = True
      run.font.color.rgb = PRIMARY_COLOR
    elif stripped.startswith("## "):
      p = doc.add_paragraph()
      p.paragraph_format.space_before = Pt(14)
      p.paragraph_format.space_after = Pt(4)
      run = p.add_run(stripped[3:])
      run.font.name = "Segoe UI"
      run.font.size = Pt(12.5)
      run.font.bold = True
      run.font.color.rgb = PRIMARY_COLOR
    elif stripped.startswith("### "):
      p = doc.add_paragraph()
      p.paragraph_format.space_before = Pt(10)
      p.paragraph_format.space_after = Pt(2)
      run = p.add_run(stripped[4:])
      run.font.name = "Segoe UI"
      run.font.size = Pt(11)
      run.font.bold = True
      run.font.color.rgb = SECONDARY_COLOR
    # Bullet points / Acceptance Criteria check items
    elif stripped.startswith("- [ ] ") or stripped.startswith("- [x] "):
      p = doc.add_paragraph(style="List Bullet")
      p.paragraph_format.space_after = Pt(3)
      check = "[✓] " if "[x]" in stripped else "[ ] "
      run = p.add_run(check)
      run.font.bold = True
      run.font.color.rgb = ACCENT_GREEN if "[x]" in stripped else PRIMARY_COLOR
      _add_formatted_runs(p, stripped[6:])
    elif stripped.startswith("- ") or stripped.startswith("* ") or stripped.startswith("• "):
      p = doc.add_paragraph(style="List Bullet")
      p.paragraph_format.space_after = Pt(3)
      _add_formatted_runs(p, stripped[2:])
    elif re.match(r"^\d+\.\s", stripped):
      remaining = re.sub(r"^\d+\.\s*", "", stripped)
      is_section_title = (
        len(remaining) >= 3
        and len(remaining) <= 80
        and not remaining.endswith((".", ",", ";", ":"))
        and (remaining == remaining.upper() or remaining[0].isupper())
        and " " in remaining
      )
      if is_section_title:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(stripped)
        run.font.name = "Segoe UI"
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = PRIMARY_COLOR
      else:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        text = re.sub(r"^\d+\.\s", "", stripped)
        _add_formatted_runs(p, text)
    # Blockquotes / Notes
    elif stripped.startswith("> "):
      p = doc.add_paragraph()
      p.paragraph_format.left_indent = Inches(0.3)
      p.paragraph_format.space_before = Pt(4)
      p.paragraph_format.space_after = Pt(4)
      _add_formatted_runs(p, stripped[2:], italic=True, color=RGBColor(71, 85, 105))
    # Normal Paragraph
    else:
      p = doc.add_paragraph()
      p.paragraph_format.space_after = Pt(6)
      p.paragraph_format.line_spacing = 1.15
      _add_formatted_runs(p, stripped)

  if in_table and table_rows:
    _render_table(doc, table_rows)

  buffer = io.BytesIO()
  doc.save(buffer)
  return buffer.getvalue()


def _add_formatted_runs(paragraph, text: str, italic: bool = False, color: RGBColor | None = None):
  parts = re.split(r"(\*\*.*?\*\*)", text)
  for part in parts:
    if part.startswith("**") and part.endswith("**"):
      run = paragraph.add_run(part[2:-2])
      run.font.bold = True
    else:
      run = paragraph.add_run(part)
      if italic:
        run.font.italic = True
    run.font.name = "Segoe UI"
    run.font.size = Pt(10.5)
    run.font.color.rgb = color if color else TEXT_COLOR


def _render_table(doc: Document, rows: list[list[str]]):
  if not rows:
    return
  col_count = max(len(r) for r in rows)
  table = doc.add_table(rows=len(rows), cols=col_count)
  table.alignment = WD_TABLE_ALIGNMENT.CENTER
  table.style = "Table Grid"

  for r_idx, row_data in enumerate(rows):
    row = table.rows[r_idx]
    is_header = (r_idx == 0)
    for c_idx, cell_value in enumerate(row_data):
      if c_idx < len(row.cells):
        cell = row.cells[c_idx]
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        if is_header:
          set_cell_background(cell, BG_HEADER_HEX)
          run = p.add_run(cell_value)
          run.font.name = "Segoe UI"
          run.font.bold = True
          run.font.size = Pt(10)
          run.font.color.rgb = RGBColor(255, 255, 255)
        else:
          if r_idx % 2 == 1:
            set_cell_background(cell, BG_ALT_ROW_HEX)
          run = p.add_run(cell_value)
          run.font.name = "Segoe UI"
          run.font.size = Pt(9.5)
          run.font.color.rgb = TEXT_COLOR
