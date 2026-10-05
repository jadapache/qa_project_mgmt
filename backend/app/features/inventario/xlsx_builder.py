"""Builder para exportación del Inventario de Requerimientos a formato Excel (.xlsx)."""

from __future__ import annotations

import io
import re
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

PRIMARY_HEX = "002777"       # Corporate Blue
SECONDARY_HEX = "004497"     # Navy Accent
HEADER_FILL_HEX = "002777"
ALT_ROW_HEX = "F0F4FF"
SECTION_FILL_HEX = "E0E7FF"
HEADER_FONT_COLOR = "FFFFFF"


def create_inventario_xlsx(
  markdown_content: str = "",
  title: str = "Inventario de Requerimientos",
  rows: list[dict[str, Any]] | None = None,
) -> bytes:
  """Construye un archivo .xlsx con el Inventario de Requerimientos y Contexto.

  Puede parsear las tablas y secciones de markdown o usar filas estructuradas.
  """
  wb = openpyxl.Workbook()

  header_fill = PatternFill(start_color=HEADER_FILL_HEX, end_color=HEADER_FILL_HEX, fill_type="solid")
  header_font = Font(bold=True, color=HEADER_FONT_COLOR, name="Segoe UI", size=11)
  section_fill = PatternFill(start_color=SECTION_FILL_HEX, end_color=SECTION_FILL_HEX, fill_type="solid")
  section_font = Font(bold=True, color="002777", name="Segoe UI", size=11)
  alt_fill = PatternFill(start_color=ALT_ROW_HEX, end_color=ALT_ROW_HEX, fill_type="solid")
  center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
  left_align = Alignment(horizontal="left", vertical="center", wrap_text=True)
  thin_border = Border(
    left=Side(style="thin", color="CBD5E1"),
    right=Side(style="thin", color="CBD5E1"),
    top=Side(style="thin", color="CBD5E1"),
    bottom=Side(style="thin", color="CBD5E1"),
  )

  # 1. Main Sheet: Inventario de Requerimientos
  ws = wb.active
  ws.title = "Requerimientos"

  # Parse markdown tables if markdown_content is given
  parsed_tables = _extract_tables_from_markdown(markdown_content) if markdown_content else []

  req_table = next((t for t in parsed_tables if any("requerimiento" in h.lower() or "id" in h.lower() for h in t["headers"])), None)

  if req_table and req_table["rows"]:
    headers = req_table["headers"]
    for col_idx, header in enumerate(headers, start=1):
      cell = ws.cell(row=1, column=col_idx, value=header)
      cell.font = header_font
      cell.fill = header_fill
      cell.alignment = center_align
      cell.border = thin_border

    for row_idx, r_data in enumerate(req_table["rows"], start=2):
      fill = alt_fill if row_idx % 2 == 0 else PatternFill()
      for col_idx, val in enumerate(r_data, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.fill = fill
        cell.alignment = left_align
        cell.border = thin_border
        cell.font = Font(name="Segoe UI", size=10)
  elif rows:
    headers = ["ID", "Requerimiento", "Tipo", "Prioridad", "Módulo", "Fuente", "Estado"]
    for col_idx, header in enumerate(headers, start=1):
      cell = ws.cell(row=1, column=col_idx, value=header)
      cell.font = header_font
      cell.fill = header_fill
      cell.alignment = center_align
      cell.border = thin_border

    for row_idx, row_data in enumerate(rows, start=2):
      fill = alt_fill if row_idx % 2 == 0 else PatternFill()
      for col_idx, key in enumerate(["id", "requirement", "type", "priority", "module", "source", "status"], start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=row_data.get(key, ""))
        cell.fill = fill
        cell.alignment = left_align
        cell.border = thin_border
        cell.font = Font(name="Segoe UI", size=10)
  else:
    # Default initial template structure
    headers = ["ID", "Requerimiento de Alto Nivel", "Tipo", "Prioridad", "Módulo Relacionado", "Fuente", "Estado"]
    default_rows = [
      ["INV-001", "Gestión de auditoría y levantamiento temprano", "Funcional", "Alta", "Módulo Funcional", "Minuta Kickoff", "Aprobado"],
      ["INV-002", "Exportación automática a Microsoft Excel (.xlsx)", "Funcional", "Alta", "Exportador", "Especificación", "Aprobado"],
      ["INV-003", "Indexación en base de conocimiento RAG", "Técnico", "Media", "Biblioteca", "Sesión Técnica", "Propuesto"],
    ]
    for col_idx, header in enumerate(headers, start=1):
      cell = ws.cell(row=1, column=col_idx, value=header)
      cell.font = header_font
      cell.fill = header_fill
      cell.alignment = center_align
      cell.border = thin_border

    for row_idx, r_data in enumerate(default_rows, start=2):
      fill = alt_fill if row_idx % 2 == 0 else PatternFill()
      for col_idx, val in enumerate(r_data, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=val)
        cell.fill = fill
        cell.alignment = left_align
        cell.border = thin_border
        cell.font = Font(name="Segoe UI", size=10)

  ws.row_dimensions[1].height = 26
  ws.freeze_panes = "A2"

  # Auto width
  for col in ws.columns:
    max_len = max(len(str(cell.value or "")) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

  # 2. Secondary Sheet: Stakeholders y Contexto
  stakeholders_table = next((t for t in parsed_tables if any("stakeholder" in h.lower() or "rol" in h.lower() for h in t["headers"])), None)
  if stakeholders_table or markdown_content:
    ws2 = wb.create_sheet(title="Stakeholders y Riesgos")
    row_cursor = 1

    if stakeholders_table and stakeholders_table["rows"]:
      ws2.cell(row=row_cursor, column=1, value="MATRIZ DE STAKEHOLDERS").font = section_font
      ws2.row_dimensions[row_cursor].height = 22
      row_cursor += 1

      for col_idx, h in enumerate(stakeholders_table["headers"], start=1):
        cell = ws2.cell(row=row_cursor, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
      row_cursor += 1

      for r_data in stakeholders_table["rows"]:
        for col_idx, val in enumerate(r_data, start=1):
          cell = ws2.cell(row=row_cursor, column=col_idx, value=val)
          cell.alignment = left_align
          cell.border = thin_border
          cell.font = Font(name="Segoe UI", size=10)
        row_cursor += 1
      row_cursor += 2

    # Riesgos Table if present
    riesgos_table = next((t for t in parsed_tables if any("riesgo" in h.lower() or "mitigación" in h.lower() or "impacto" in h.lower() for h in t["headers"])), None)
    if riesgos_table and riesgos_table["rows"]:
      ws2.cell(row=row_cursor, column=1, value="MATRIZ DE RIESGOS Y SUPUESTOS").font = section_font
      ws2.row_dimensions[row_cursor].height = 22
      row_cursor += 1

      for col_idx, h in enumerate(riesgos_table["headers"], start=1):
        cell = ws2.cell(row=row_cursor, column=col_idx, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align
        cell.border = thin_border
      row_cursor += 1

      for r_data in riesgos_table["rows"]:
        for col_idx, val in enumerate(r_data, start=1):
          cell = ws2.cell(row=row_cursor, column=col_idx, value=val)
          cell.alignment = left_align
          cell.border = thin_border
          cell.font = Font(name="Segoe UI", size=10)
        row_cursor += 1

    for col in ws2.columns:
      max_len = max(len(str(cell.value or "")) for cell in col)
      col_letter = get_column_letter(col[0].column)
      ws2.column_dimensions[col_letter].width = max(max_len + 4, 14)

  buffer = io.BytesIO()
  wb.save(buffer)
  buffer.seek(0)
  return buffer.read()


def _extract_tables_from_markdown(markdown_content: str) -> list[dict[str, Any]]:
  tables = []
  lines = markdown_content.splitlines()
  current_table_rows: list[list[str]] = []
  in_table = False

  for line in lines:
    stripped = line.strip()
    if stripped.startswith("|") and stripped.endswith("|"):
      if re.match(r"^\|[\s\-:|]+\|$", stripped):
        continue
      cells = [c.strip() for c in stripped.split("|")[1:-1]]
      current_table_rows.append(cells)
      in_table = True
    elif in_table:
      if len(current_table_rows) >= 2:
        tables.append({
          "headers": current_table_rows[0],
          "rows": current_table_rows[1:],
        })
      current_table_rows = []
      in_table = False

  if in_table and len(current_table_rows) >= 2:
    tables.append({
      "headers": current_table_rows[0],
      "rows": current_table_rows[1:],
    })

  return tables
